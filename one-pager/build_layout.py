"""Build the designed one-pager as an editable Letter-size slide file.

Run:  python3 one-pager/build_layout.py
Out:  one-pager/368-Hellhound-OnePager-Layout.pptx
      Open in Canva (Upload) or Google Slides; every box and line is editable.
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
OUT = HERE / "368-Hellhound-OnePager-Layout.pptx"

PAPER = "FAF9F6"
CARD = "FFFFFF"
INK = "171717"
MUTED = "6B6B6B"
FAINT = "A3A3A3"
EDGE = "ECEAE5"
RULE = "F0EEEA"
PILL = "EEF6FB"
BLUE = "0096CF"
FONT = "Inter"

W, H, M = 8.5, 11, 0.6      # page and side margin, inches
INNER = W - 2 * M

# Sample content. Replace per client (numbers come from the Sheet's Inputs tab).
CLIENT = "Acme Corp"
DATE = "Nov 12, 2026"
EVENT = "Q4 Leadership Town Hall"
EVENT_TYPE = "Executive Zoom Call"
CONTACT = ["Your name", "you@company.com", "(555) 555-5555"]
PROOF = [("1,000+", "events produced"), ("24 hrs", "of testing before every show"),
         ("2×", "backup on every critical part")]
PACKAGES = [
    ("Sound", "Wireless mics, spares, live mixing", 2000),
    ("Video & Computers", "Backup computer, cameras, wired internet", 2500),
    ("Screens", "Displays for the room", 900),
    ("Crew & Rehearsal", "Lead engineer, call operator, rehearsal day", 3900),
]
STANDARD = ["Backup computer, running live", "Wired internet with cellular backup",
            "A mic for every speaker", "Lead engineer on site",
            "Full rehearsal the day before", "Delivery, setup, and teardown"]
LINES = [  # (category, item, description, qty, days, rate)
    ("Sound", "Executive Mic Package", "Wireless mic per speaker, plus spares", 1, 1, 1200),
    ("Sound", "Live Audio Mixing", "Live mixing for room and remote audio", 1, 1, 800),
    ("Video & Computers", "Show Computer Package", "Primary and backup show computers", 1, 1, 900),
    ("Video & Computers", "Camera & Lighting Package", "Cameras and lighting for speakers", 1, 1, 1100),
    ("Video & Computers", "Dedicated Internet", "Wired line with cellular backup", 1, 1, 500),
    ("Screens", "Display Package", "Display, stand, and cabling", 2, 1, 450),
    ("Crew & Rehearsal", "Lead Engineer", "Runs the show", 1, 1, 1400),
    ("Crew & Rehearsal", "Call Operator", "Runs Zoom: guests, shares, recording", 1, 1, 1000),
    ("Crew & Rehearsal", "Setup & Rehearsal Day", "Build and full test, day before", 1, 1, 1500),
]
TAX = 0.0
NOTES = "Load-in begins 7:00 AM. Final rehearsal with speakers at 2:00 PM the day before."


def rgb(h):
    return RGBColor.from_string(h)


def money(n):
    return f"${n:,.0f}"


def effects(shape, shadow):
    """Very faint drop shadow (6% black, 10pt blur, 1.5pt down), or none at all.
    An explicit empty list also stops the theme adding its default shadow."""
    ref = shape._element.find(".//" + qn("a:effectRef"))
    if ref is not None:
        ref.set("idx", "0")
    sp_pr = shape._element.spPr
    for old in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(old)
    eff = etree.SubElement(sp_pr, qn("a:effectLst"))
    if shadow:
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(127000), dist=str(19050),
                              dir="5400000", algn="t", rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val="6000")


def card(slide, x, y, w, h, radius=0.14, fill=CARD, edge=EDGE, shadow=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if edge:
        s.line.color.rgb = rgb(edge)
        s.line.width = Pt(0.5)
    else:
        s.line.fill.background()
    effects(s, shadow)
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, value, size=10, bold=False, color=INK, align="left", anchor="top"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    return tb


def rule(slide, x, y, w, color=RULE):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(0.75))
    ln.fill.solid()
    ln.fill.fore_color.rgb = rgb(color)
    ln.line.fill.background()
    effects(ln, False)


def new_page(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    text(s, M, 0.5, 4, 0.2, "368 EVENT MANAGEMENT  ×  HELLHOUND AUDIO", 7.5, True, MUTED)
    text(s, W - M - 4, 0.5, 4, 0.2, f"Prepared for {CLIENT}  ·  {DATE}", 7.5, False, MUTED, "right")
    return s


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)

# ---------------------------------------------------------------- page 1
s = new_page(prs)
pill = card(s, M, 1.1, 1.55, 0.28, radius=0.14, fill=PILL, edge=None, shadow=False)
text(s, M, 1.1, 1.55, 0.28, EVENT_TYPE, 8, True, BLUE, "center", "middle")
text(s, M, 1.5, INNER, 0.6, "Your call just works.", 34, True)

gap = 0.2
cw = (INNER - 2 * gap) / 3
for i, (num, label) in enumerate(PROOF):
    x = M + i * (cw + gap)
    card(s, x, 2.4, cw, 0.95)
    text(s, x + 0.22, 2.55, cw - 0.44, 0.4, num, 22, True)
    text(s, x + 0.22, 2.98, cw - 0.44, 0.22, label, 8.5, False, MUTED)

text(s, M, 3.65, INNER, 0.3, "What’s included", 13, True)
cw = (INNER - gap) / 2
for i, (name, spec, price) in enumerate(PACKAGES):
    x, y = M + (i % 2) * (cw + gap), 4.05 + (i // 2) * (1.2 + gap)
    card(s, x, y, cw, 1.2)
    text(s, x + 0.25, y + 0.2, cw - 0.5, 0.25, name, 12, True)
    text(s, x + 0.25, y + 0.47, cw - 0.5, 0.22, spec, 8.5, False, MUTED)
    text(s, x + 0.25, y + 0.76, cw - 0.5, 0.3, money(price), 16, True)

y = 6.85
text(s, M, y, INNER, 0.3, "Standard on every show", 13, True)
card(s, M, y + 0.4, INNER, 1.1)
for i, item in enumerate(STANDARD):
    cx = M + 0.25 + (i // 3) * (INNER / 2)
    cy = y + 0.58 + (i % 3) * 0.27
    text(s, cx, cy, 0.2, 0.22, "✓", 9, True, BLUE)
    text(s, cx + 0.22, cy, INNER / 2 - 0.6, 0.22, item, 9.5)

y = 8.6
card(s, M, y, INNER, 1.0)
text(s, M + 0.25, y + 0.2, 2, 0.2, "Total", 9, False, MUTED)
subtotal = sum(q * d * r for _, _, _, q, d, r in LINES)
total = subtotal * (1 + TAX)
text(s, M + 0.25, y + 0.4, 3.5, 0.45, money(total), 28, True)
text(s, W - M - 3.25, y + 0.25, 3, 0.55, "\n".join(CONTACT), 9, False, MUTED, "right", "middle")

s.shapes.add_picture(str(HERE / "assets" / "368em-logo.png"), Inches(M), Inches(9.9), Inches(0.62), Inches(0.62))
s.shapes.add_picture(str(HERE / "assets" / "hellhound-logo.png"), Inches(W - M - 1.95), Inches(9.92),
                     Inches(1.95), Inches(0.6))

# ---------------------------------------------------------------- page 2
s = new_page(prs)
text(s, M, 1.05, INNER, 0.45, "Line items", 24, True)
text(s, M, 1.55, INNER, 0.22, EVENT, 10, False, MUTED)

cols = [("ITEM", 1.95, "left"), ("DESCRIPTION", 2.4, "left"), ("QTY", 0.45, "center"),
        ("DAYS", 0.45, "center"), ("RATE", 0.75, "right"), ("SUBTOTAL", 0.8, "right")]
cats = list(dict.fromkeys(c for c, *_ in LINES))
row_h = 0.3
top = 2.0
card_h = 0.25 + 0.3 + (len(cats) + len(LINES)) * row_h + 0.2
card(s, M, top, INNER, card_h)
x0 = M + 0.25
x = x0
for label, w, al in cols:
    text(s, x, top + 0.25, w, 0.2, label, 7, True, FAINT, al)
    x += w
y = top + 0.55
for cat in cats:
    text(s, x0, y + 0.08, 3, 0.2, cat.upper(), 7.5, True, BLUE)
    y += row_h
    items = [l for l in LINES if l[0] == cat]
    for j, (_, item, desc, q, d, r) in enumerate(items):
        vals = [item, desc, str(q), str(d), money(r), money(q * d * r)]
        styles = [(9.5, True, INK), (8.5, False, MUTED), (9.5, False, INK), (9.5, False, INK),
                  (9.5, False, MUTED), (9.5, True, INK)]
        x = x0
        for (label, w, al), v, (sz, b, c) in zip(cols, vals, styles):
            text(s, x, y, w, row_h, v, sz, b, c, al, "middle")
            x += w
        if j < len(items) - 1:
            rule(s, x0, y + row_h, INNER - 0.5)
        y += row_h

y = top + card_h + 0.25
tw = 3.0
card(s, W - M - tw, y, tw, 1.15)
for k, (label, val, sz, b) in enumerate([("Subtotal", money(subtotal), 9.5, False),
                                         (f"Tax ({TAX:.1%})", money(subtotal * TAX), 9.5, False),
                                         ("Total", money(total), 13, True)]):
    yy = y + 0.18 + k * 0.28
    text(s, W - M - tw + 0.25, yy, 1.4, 0.26, label, sz, b, INK if b else MUTED, "left", "middle")
    text(s, W - M - 1.6, yy, 1.35, 0.26, val, sz, b, INK, "right", "middle")
rule(s, W - M - tw + 0.25, y + 0.73, tw - 0.5, EDGE)

y += 1.45
card(s, M, y, INNER, 0.95)
text(s, M + 0.25, y + 0.2, 2, 0.22, "Notes", 10, True)
text(s, M + 0.25, y + 0.45, INNER - 0.5, 0.35, NOTES, 9.5, False, MUTED)

text(s, M, 10.3, INNER, 0.2, "368 Event Management  ×  Hellhound Audio  ·  hellhoundaudio.com",
     7.5, False, FAINT, "center")

prs.save(OUT)
print("wrote", OUT)
