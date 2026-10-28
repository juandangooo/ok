"""Build the one-page quote as an editable Letter-size file for Canva / Google Slides.

Run:  python3 one-pager/quote/editable.py one-pager/quote/helix-2026-10-22.json
Out:  <name>-editable.pptx  (Canva: Create a design › Import file)
Every text line, number, and rule is its own element.
"""
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK, MUTED, FAINT, LINE, BLUE, RED = "171717", "6B6B6B", "A3A3A3", "E7E5E0", "0096CF", "E8382F"
FONT = "Inter"
W, H, L = 8.5, 11, 0.7
R = W - L
INNER = R - L

src = Path(sys.argv[1])
q = json.loads(src.read_text())
o = q["onepage"]


def money(n):
    return f"${n:,.0f}"


totals = {g["name"]: sum(a * b * c for _, _, a, b, c in g["items"]) for g in q["groups"]}
kinds = {g["name"]: g["kind"] for g in q["groups"]}
for src_name, dst in o.get("fold", {}).items():  # show one group's cost inside another
    totals[dst] += totals.pop(src_name)
    kinds.pop(src_name)
rental = sum(v for k, v in totals.items() if kinds[k] == "rental")
service = sum(v for k, v in totals.items() if kinds[k] == "service")
tax = (rental + service) * q["tax"]
total = rental + service + tax

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")


def text(x, y, w, h, value, size, bold=False, color=INK, align="left", anchor="top", name=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


def rule(x, y, w, color=LINE, weight=0.75):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(weight))
    ln.name = "Rule"
    ln.fill.solid()
    ln.fill.fore_color.rgb = RGBColor.from_string(color)
    ln.line.fill.background()
    ref = ln._element.find(".//" + qn("a:effectRef"))
    if ref is not None:
        ref.set("idx", "0")
    etree.SubElement(ln._element.spPr, qn("a:effectLst"))


text(L, 0.6, 4, 0.18, "368 EVENT MANAGEMENT  ×  HELLHOUND AUDIO", 7.5, True, MUTED)
text(R - 4, 0.6, 4, 0.18, f"Prepared for {q['client']}  ·  Quote {q['quote_number']}", 7.5, False, MUTED, "right")
text(L, 0.86, 5, 0.2, o["kicker"].upper(), 8.5, True, BLUE, name="Kicker")
text(L, 1.0, 5.5, 1.3, o["headline"], 38, True, name="Headline")

text(L, 2.55, 4, 0.3, "What’s included", 14, True)
y = 2.9
for name, desc in o["included"].items():
    rule(L, y, INNER)
    text(L, y + 0.13, 5.2, 0.22, name, 11, True, name=f"{name} name")
    text(L, y + 0.34, 5.2, 0.2, desc, 8.5, False, MUTED, name=f"{name} description")
    text(R - 1.6, y, 1.6, 0.575, money(totals[name]), 14, True, align="right", anchor="middle",
         name=f"{name} price")
    y += 0.575

text(L, 5.15, 4, 0.3, o["also_title"], 14, True)
half = (INNER - 0.25) / 2
for i, item in enumerate(o["also"]):
    x = L + (i // 3) * (half + 0.25)
    yy = 5.55 + (i % 3) * 0.29
    text(x, yy, half, 0.22, f"✓   {item}", 9, name=f"Also {i + 1}")
    rule(x, yy + 0.27, half)

rule(L, 7.0, INNER, INK, 1.25)
text(L, 7.11, 2, 0.16, "TOTAL", 7.5, True, RED)
text(L, 7.25, 4, 0.55, money(total), 32, True, name="Total")
for k, (label, val) in enumerate([("Equipment", rental), ("Crew", service), ("Tax", tax)]):
    yy = 7.17 + k * 0.19
    text(R - 2.6, yy, 1.8, 0.17, label, 8.5, False, MUTED)
    text(R - 0.8, yy, 0.8, 0.17, money(val), 8.5, False, MUTED, "right", name=f"{label} amount")

text(L, 8.1, INNER, 0.2, o["contact"], 9, False, MUTED, name="Contact")
text(L, 10.25, INNER, 0.18,
     f"Quote {q['quote_number']}  ·  Issued {q['quote_date']}  ·  Valid until {q['valid_until']}",
     7.5, False, FAINT, name="Validity")

out = src.with_name(src.stem + "-editable.pptx")
prs.save(out)
print("wrote", out)
