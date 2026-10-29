"""Build the step-by-step guide deck (opens in Google Slides).

Run:  python3 one-pager/guide/mocks.py && node one-pager/guide/capture.js && python3 one-pager/guide/build_deck.py
Out:  one-pager/guide/How-to-Make-a-Quote.pptx
"""
import json
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
MOCKS = HERE / "mocks"
OUT = HERE / "How-to-Make-a-Quote.pptx"
BOXES = json.loads((MOCKS / "boxes.json").read_text())

INK, MUTED, RED, BLUE, SOFT = "171717", "6B6B6B", "E8382F", "0096CF", "F5F5F7"
FONT = "Inter"
SW, SH = 13.333, 7.5
IMG_X, IMG_Y, IMG_W, IMG_H = 4.55, 0.85, 8.35, 5.9   # screenshot area on step slides
CROPS = HERE / "mocks" / "crops"
CROPS.mkdir(exist_ok=True)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)


def rgb(h):
    return RGBColor.from_string(h)


def no_effects(shape):
    ref = shape._element.find(".//" + qn("a:effectRef"))
    if ref is not None:
        ref.set("idx", "0")
    etree.SubElement(shape._element.spPr, qn("a:effectLst"))


def text(slide, x, y, w, h, runs, size=20, color=INK, align="left", anchor="top", spacing=1.1):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, bold) pieces."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
        p.line_spacing = spacing
        if i:
            p.space_before = Pt(size * 0.6)
        for piece, bold in ([(para, False)] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = piece
            r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
            r.font.color.rgb = rgb(color)
    return tb


def box(slide, x, y, w, h, fill=None, line=None, weight=0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(weight)
    else:
        s.line.fill.background()
    no_effects(s)
    return s


def new_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb("FFFFFF")
    return s


def screenshot(slide, name, crop):
    """Place a zoomed-in part of a mock (crop = x, y, w, h in mock pixels); return its placement."""
    x, y, w, h = crop
    im = Image.open(MOCKS / f"{name}.png")        # captured at 2x
    path = CROPS / f"{name}-{x}-{y}-{w}-{h}.png"
    im.crop((x * 2, y * 2, (x + w) * 2, (y + h) * 2)).save(path)
    scale = min(IMG_W / w, IMG_H / h)              # inches per mock pixel
    dw, dh = w * scale, h * scale
    slide.shapes.add_picture(str(path), Inches(IMG_X), Inches(IMG_Y), Inches(dw), Inches(dh))
    box(slide, IMG_X, IMG_Y, dw, dh, line="DADCE0", weight=0.75, radius=0.06)
    return {"crop": crop, "scale": scale, "w": dw, "h": dh}


def circle(slide, name, place, target, label=None, label_side="below", pad=10):
    x, y, w, h = BOXES[name][target]
    k, (ox, oy, _, _) = place["scale"], place["crop"]
    cx, cy = IMG_X + (x - ox - pad) * k, IMG_Y + (y - oy - pad) * k
    cw, ch = (w + 2 * pad) * k, (h + 2 * pad) * k
    box(slide, cx, cy, cw, ch, line=RED, weight=4.5, shape=MSO_SHAPE.OVAL)
    if label:
        lw = 2.7
        if label_side in ("left", "right"):
            lx = cx - lw - 0.12 if label_side == "left" else cx + cw + 0.12
            ly = cy + ch / 2 - 0.23
        else:
            lx = cx + cw / 2 - lw / 2
            ly = cy + ch + 0.08 if label_side == "below" else cy - 0.55
        lx = min(max(lx, IMG_X + 0.1), IMG_X + place["w"] - lw - 0.1)
        box(slide, lx, ly, lw, 0.46, fill=RED, radius=0.23)
        text(slide, lx, ly, lw, 0.46, [[(label, True)]], 16, "FFFFFF", "center", "middle")


FORM = (0, 180, 800, 560)      # labels and blue cells
BUTTONS = (560, 180, 760, 540)  # hints and the two buttons


def step(badge, title, lines, shot, crop, circles, size=20):
    s = new_slide()
    label = f"STEP {badge}" if isinstance(badge, int) else badge
    bw = 0.32 + 0.13 * len(label)
    box(s, 0.6, 0.85, bw, 0.5, fill=BLUE, radius=0.25)
    text(s, 0.6, 0.85, bw, 0.5, [[(label, True)]], 14, "FFFFFF", "center", "middle")
    text(s, 0.6, 1.6, 3.7, 1.6, [[(title, True)]], 30, spacing=1.05)
    text(s, 0.6, 3.4, 3.6, 3.9, lines, size, MUTED, spacing=1.2)
    place = screenshot(s, shot, crop)
    for c in circles:
        circle(s, shot, place, *c)
    return s


# 1 · title
s = new_slide()
text(s, 0.9, 2.3, 11, 1.2, [[("How to make a quote", True)]], 54)
text(s, 0.9, 3.55, 11, 0.6, "Helix package quotes  ·  pick a package, fill in the blue cells, click Make PDF", 22, MUTED)
text(s, 0.9, 6.6, 11, 0.4, "368 Event Management  ×  Hellhound Audio", 13, MUTED)

# 2 · what it does
s = new_slide()
text(s, 0.9, 0.9, 11, 0.8, [[("What it does", True)]], 36)
cards = [("1", "You type the details", "Client, date, package"),
         ("2", "You click Make PDF", "One button"),
         ("3", "You get the quote", "A clean one-page PDF, saved in Drive")]
for i, (num, head, sub) in enumerate(cards):
    x = 0.9 + i * 3.95
    box(s, x, 2.3, 3.55, 3.0, fill=SOFT, radius=0.2)
    text(s, x + 0.4, 2.65, 2.8, 0.7, [[(num, True)]], 40, BLUE)
    text(s, x + 0.4, 3.6, 2.9, 0.9, [[(head, True)]], 22)
    text(s, x + 0.4, 4.45, 2.9, 0.7, sub, 16, MUTED)
    if i < 2:
        text(s, x + 3.58, 3.4, 0.4, 0.6, [[("→", True)]], 28, MUTED, "center")
text(s, 0.9, 5.85, 11.5, 0.6, [[("No math. No line items. ", True), ("Prices come from the Packages tab.", False)]], 18, MUTED)

# 3-8 · the steps
step(1, "Open the sheet. Click the Quote tab.",
     [[("Open ", False), ("Helix Package Quotes", True), (" in Google Drive.", False)],
      [("The tabs are at the ", False), ("bottom", True), (" of the screen.", False)]],
     "quote", (0, 340, 800, 560), [("tab-quote", "Click here", "above")])
step(2, "Click Start new quote.",
     ["It clears the old client and date.",
      [("The quote number goes up by one ", False), ("on its own", True), (".", False)]],
     "quote", BUTTONS, [("btn-new", "Click here", "left")])
step(3, "Pick the package.",
     [[("Click the ", False), ("Package", True), (" cell.", False)],
      [("Pick from the list. ", False), ("Don't type it.", True)]],
     "quote", FORM, [("cell-package", "Pick from the list", "right")])
step(4, "Fill in the blue cells.",
     [[("Only type in the ", False), ("blue", True), (" cells.", False)],
      "The gray words next to each one tell you what goes there.",
      [("Add-ons", True), (" at the bottom are optional.", False)]],
     "quote-empty", FORM, [("cell-client", None), ("cell-date", "Type these in", "right")])
step(5, "Click Make PDF.",
     ["Wait a few seconds.",
      [("If something is missing, it ", False), ("tells you what to fix", True), (".", False)]],
     "quote", BUTTONS, [("btn-make", "Click here", "left")])
step(6, "Click Open PDF.",
     [[("That's the quote. ", False), ("Send it to the client.", True)],
      [("Need to change a word? Click ", False), ("Edit as Google Doc", True), (".", False)]],
     "done", (380, 160, 900, 620), [("open-pdf", "Click here")])

# 9 · the result
s = new_slide()
text(s, 0.6, 1.05, 3.8, 1.2, [[("Your PDF looks like this.", True)]], 30, spacing=1.05)
text(s, 0.6, 2.55, 3.6, 3.5,
     ["One page. Letter size.",
      [("The package, the price, and ", False), ("nothing extra", True), (".", False)],
      "Ready to print or email."], 18, MUTED, spacing=1.2)
pic_h = 6.4
s.shapes.add_picture(str(MOCKS / "onepager.png"), Inches(6.4), Inches(0.55), height=Inches(pic_h))
box(s, 6.4, 0.55, pic_h * 8.5 / 11, pic_h, line="DADCE0", weight=0.75, radius=0.04)

# 10 · where files go
step("WHERE FILES GO", "Where your files go.",
         [[("Every PDF is saved in the ", False), ("Helix Quotes", True), (" folder in Google Drive.", False)],
          [("The ", False), ("Log", True), (" tab lists every quote, with links.", False)]],
         "log", (0, 180, 800, 560), [("log-row", "Every quote, saved")])

# 11 · fix these first
step("IF IT STOPS", "It says “Fix these first.”",
     ["That means a blue cell is empty or wrong.",
      [("Read the list. ", True), ("Fix those cells. Click ", False), ("Make PDF", True), (" again.", False)]],
     "fix", (380, 150, 900, 620), [("alert", "Read this list", "above")])

# 12 · prices
step("PRICES", "Change a price.",
     [[("Go to the ", False), ("Packages", True), (" tab. Type the new price.", False)],
      [("Red rows", True), (" have fake prices. Put the real price in, then delete the word PLACEHOLDER.", False)]],
     "packages", (500, 180, 960, 640), [("ph-price", None), ("ph-status", "Fake price: fix these")])

# 13 · rules
s = new_slide()
text(s, 0.9, 0.9, 11, 0.8, [[("Three rules", True)]], 36)
rules = [("Only type in blue cells.", "Everything else fills itself in."),
         ("Red rows are fake prices.", "Fix them before quoting that package."),
         ("Use a computer.", "The buttons don't work in the phone app. Use the Quotes menu there.")]
for i, (head, sub) in enumerate(rules):
    y = 2.15 + i * 1.5
    box(s, 0.9, y, 11.5, 1.25, fill=SOFT, radius=0.18)
    text(s, 1.3, y + 0.22, 0.6, 0.8, [[(str(i + 1), True)]], 30, BLUE)
    text(s, 2.1, y + 0.2, 10, 0.5, [[(head, True)]], 22)
    text(s, 2.1, y + 0.68, 10, 0.45, sub, 16, MUTED)

# 14 · setup
step("SETUP", "One-time setup.",
     [[("Only do this once.", True)],
      "1. Make a blank Google Sheet.",
      [("2. ", False), ("Extensions › Apps Script", True), (". Paste the script. Save.", False)],
      "3. Reload the sheet.",
      [("4. ", False), ("Quotes › Set up sheet", True), (". Click Allow.", False)],
      "5. Share it with the sales team."],
     "menu", (300, 0, 820, 580), [("menu-setup", "Click once")], size=17)

prs.save(OUT)
print("wrote", OUT, len(prs.slides), "slides")
