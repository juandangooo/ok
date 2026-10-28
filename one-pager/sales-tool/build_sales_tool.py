"""Build the Helix package quote tool (Google Sheets / Excel workbook).

Run:  python3 one-pager/sales-tool/build_sales_tool.py
Out:  one-pager/sales-tool/Helix-Package-Quotes.xlsx

Tabs: Start · Quote (sales enters details) · One-Pager (prints as the PDF) · Packages (library)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

HERE = Path(__file__).parent
OUT = HERE / "Helix-Package-Quotes.xlsx"

INK, MUTED, FAINT, LINE = "171717", "6B6B6B", "A3A3A3", "E7E5E0"
BLUE, RED, INPUT, WARN = "0096CF", "E8382F", "EAF4FA", "FDECEA"
FONT = "Inter"
MONEY = '"$"#,##0'
SLOTS = 5      # package sections + add-ons shown on the one-pager
BULLETS = 6

# ---------------------------------------------------------------- package library
# Live Zoom numbers come from Helix quote 0625 (transport folded into Audio).
# Everything marked PLACEHOLDER needs Rick's real price before use.
PACKAGES = [
    ("Live Zoom Package", "Your Zoom\nCall Package"),
    ("Cocktail Hour Sound", "Your Cocktail\nHour Sound"),
    ("Room Rental", "Your Room\nat Helix"),
]
SECTIONS = [  # package, section, description, price per day, status
    ("Live Zoom Package", "Crew", "Senior audio technician (A1), video technician (V1)", 1200, "From quote 0625"),
    ("Live Zoom Package", "Video", "BirdDog X5 Ultra PTZ camera, Roland V-02HD MK II mixer, MacBook Air M2", 850, "From quote 0625"),
    ("Live Zoom Package", "Audio", "Focusrite Scarlett 2i2, Radial direct boxes, iPad", 600, "From quote 0625 (includes transport)"),
    ("Cocktail Hour Sound", "Crew", "Audio engineer (A1), on site for the event", 600, "PLACEHOLDER"),
    ("Cocktail Hour Sound", "Sound system", "PA speakers, digital mixer, wireless mics", 750, "PLACEHOLDER"),
    ("Cocktail Hour Sound", "Band setup", "Stage mics, monitors, and direct boxes for the band", 400, "PLACEHOLDER"),
    ("Room Rental", "Room", "Event space at Helix, 20 French St, New Brunswick", 5000, "PLACEHOLDER"),
    ("Room Rental", "Presentation AV", "Display, presentation audio, wireless mics", 900, "PLACEHOLDER"),
    ("Room Rental", "Crew", "On-site technician for the event", 600, "PLACEHOLDER"),
]
ALSO = [  # package, bullet
    ("Live Zoom Package", "Joystick camera controller"),
    ("Live Zoom Package", "SDI and HDMI cable package"),
    ("Live Zoom Package", "Remote pan, tilt, and zoom"),
    ("Live Zoom Package", "4K camera monitor for the technician"),
    ("Live Zoom Package", "PoE+ network switch"),
    ("Live Zoom Package", "Camera tripods"),
    ("Cocktail Hour Sound", "Background music playback"),
    ("Cocktail Hour Sound", "Speeches and toasts mic"),
    ("Cocktail Hour Sound", "All cabling and stands"),
    ("Room Rental", "Tables and chairs"),
    ("Room Rental", "Wi-Fi"),
    ("Room Rental", "Setup and breakdown"),
]
PKG_FIRST, PKG_LAST = 6, 25
SEC_FIRST, SEC_LAST = 6, 105
ALSO_FIRST, ALSO_LAST = 6, 155


def fill(h):
    return PatternFill("solid", start_color=h, end_color=h)


thin = Side(style="thin", color=LINE)
ink = Side(style="medium", color=INK)


def put(ws, ref, value, size=10, bold=False, color=INK, h="left", v="center", wrap=False,
        merge=None, fmt=None, italic=False):
    if merge:
        ws.merge_cells(merge)
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if fmt:
        c.number_format = fmt
    return c


def inp(ws, ref, value, fmt=None):
    c = put(ws, ref, value, 10.5, fmt=fmt)
    c.fill = fill(INPUT)
    return c


wb = Workbook()

# ---------------------------------------------------------------- Start
st = wb.active
st.title = "Start"
st.sheet_view.showGridLines = False
st.column_dimensions["A"].width = 3
st.column_dimensions["B"].width = 5
st.column_dimensions["C"].width = 92
put(st, "B2", "Helix package quotes", 22, True, merge="B2:C2")
st.row_dimensions[2].height = 34
put(st, "B3", "Pick a package, enter the details, download the one-pager as a PDF.", 11, False, MUTED, merge="B3:C3")
steps = [
    ("Quote", "Choose the package from the dropdown, then fill in the blue cells. Add-ons are optional."),
    ("One-Pager", "This is the PDF. Check it, then File › Download › PDF. Use: Current sheet · Letter · Portrait · "
                  "Fit to width · Margins: Narrow · uncheck Show gridlines."),
    ("Packages", "The package library. Change prices, wording, or add a package here once and every quote uses it. "
                 "Rows marked PLACEHOLDER need real prices before anything goes out."),
]
r = 5
for i, (name, text) in enumerate(steps, 1):
    put(st, f"B{r}", i, 14, True, BLUE, h="center", v="top")
    put(st, f"C{r}", name, 11, True)
    put(st, f"C{r+1}", text, 10.5, False, MUTED, wrap=True, v="top")
    st.row_dimensions[r + 1].height = 32
    r += 3

# ---------------------------------------------------------------- Quote (inputs)
q = wb.create_sheet("Quote")
q.sheet_view.showGridLines = False
for k, w in {"A": 22, "B": 40, "C": 50, "D": 14}.items():
    q.column_dimensions[k].width = w
put(q, "A1", "Quote", 20, True)
q.row_dimensions[1].height = 30
put(q, "A2", "Fill in the blue cells. The One-Pager tab updates on its own.", 10.5, False, MUTED)

fields = [
    (4, "Package", "Live Zoom Package", None),
    (5, "Client", "Helix", None),
    (6, "Event date", "=DATE(2026,10,22)", "mmmm d, yyyy"),
    (7, "Days", 1, "0"),
    (8, "Quote number", "0625", None),
    (9, "Issued", "=DATE(2026,10,28)", "mmm d, yyyy"),
    (10, "Valid until", "=B9+30", "mmm d, yyyy"),
    (11, "Tax rate", 0, "0.000%"),
    (12, "Contact name", "Rick", None),
    (13, "Contact email", "Rick@368eventmanagement.com", None),
    (14, "Contact phone", "(201) 893-5224", None),
]
for row, label, val, fmt in fields:
    put(q, f"A{row}", label, 10.5, False, MUTED)
    inp(q, f"B{row}", val, fmt)
dv = DataValidation(type="list", formula1=f"=Packages!$A${PKG_FIRST}:$A${PKG_LAST}", allow_blank=False)
q.add_data_validation(dv)
dv.add("B4")

put(q, "A16", "Add-ons (optional)", 11, True)
put(q, "B16", "Shown after the package sections. Price is flat, not multiplied by days.", 9.5, False, MUTED, merge="B16:D16")
for col, h in zip("ABCD", ["", "Name", "Description", "Price"]):
    put(q, f"{col}17", h, 9, True, MUTED).border = Border(bottom=thin)
for row in (18, 19):
    put(q, f"A{row}", f"Add-on {row - 17}", 10.5, False, MUTED)
    inp(q, f"B{row}", None)
    inp(q, f"C{row}", None)
    inp(q, f"D{row}", None, MONEY)

put(q, "A21", "Total", 12, True)
put(q, "B21", "='One-Pager'!B39", 16, True, fmt=MONEY)
put(q, "C21", "Same number as the One-Pager.", 9.5, False, MUTED)

# ---------------------------------------------------------------- Packages (library)
pk = wb.create_sheet("Packages")
pk.sheet_view.showGridLines = False
widths = {"A": 24, "B": 22, "C": 2, "D": 8, "E": 22, "F": 18, "G": 58, "H": 12, "I": 30,
          "J": 2, "K": 8, "L": 22, "M": 40}
for k, w in widths.items():
    pk.column_dimensions[k].width = w
put(pk, "A1", "Packages", 20, True)
pk.row_dimensions[1].height = 30
put(pk, "A2", "Edit prices and wording here. Add a package: add its name and headline on the left, then its sections "
              "and bullets using the same name.", 10, False, MUTED)
put(pk, "A3", "Red rows are PLACEHOLDER prices. Replace them before sending a quote for that package.", 10, True, RED)

for col, h in zip(["A", "B"], ["Package", "Headline (two lines)"]):
    put(pk, f"{col}5", h, 9, True, MUTED).border = Border(bottom=thin)
for col, h in zip(["D", "E", "F", "G", "H", "I"], ["Key", "Package", "Section", "Description", "Price / day", "Status"]):
    put(pk, f"{col}5", h, 9, True, MUTED).border = Border(bottom=thin)
for col, h in zip(["K", "L", "M"], ["Key", "Package", "Also included"]):
    put(pk, f"{col}5", h, 9, True, MUTED).border = Border(bottom=thin)

for i, (name, headline) in enumerate(PACKAGES):
    r = PKG_FIRST + i
    inp(pk, f"A{r}", name)
    c = inp(pk, f"B{r}", headline)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    pk.row_dimensions[r].height = 30

for r in range(SEC_FIRST, SEC_LAST + 1):
    # key = package|n, the nth section of that package; drives the one-pager lookups
    put(pk, f"D{r}", f'=IF(E{r}="","",E{r}&"|"&COUNTIF($E${SEC_FIRST}:E{r},E{r}))', 8, False, FAINT)
    pk[f"H{r}"].number_format = MONEY
for i, (pkg, sec, desc, price, status) in enumerate(SECTIONS):
    r = SEC_FIRST + i
    inp(pk, f"E{r}", pkg)
    inp(pk, f"F{r}", sec)
    inp(pk, f"G{r}", desc)
    inp(pk, f"H{r}", price, MONEY)
    put(pk, f"I{r}", status, 9, status == "PLACEHOLDER", RED if status == "PLACEHOLDER" else MUTED)
pk.conditional_formatting.add(f"E{SEC_FIRST}:I{SEC_LAST}",
                              FormulaRule(formula=[f'$I{SEC_FIRST}="PLACEHOLDER"'], fill=fill(WARN)))

for r in range(ALSO_FIRST, ALSO_LAST + 1):
    put(pk, f"K{r}", f'=IF(L{r}="","",L{r}&"|"&COUNTIF($L${ALSO_FIRST}:L{r},L{r}))', 8, False, FAINT)
for i, (pkg, bullet) in enumerate(ALSO):
    r = ALSO_FIRST + i
    inp(pk, f"L{r}", pkg)
    inp(pk, f"M{r}", bullet)
pk.freeze_panes = "A6"

# ---------------------------------------------------------------- One-Pager (prints as the PDF)
op = wb.create_sheet("One-Pager", 2)
op.sheet_view.showGridLines = False
for k, w in {"A": 2, "B": 11.75, "C": 11.75, "D": 11.75, "E": 11.75, "F": 2.5,
             "G": 11.75, "H": 11.75, "I": 11.75, "J": 11.75, "K": 2}.items():
    op.column_dimensions[k].width = w
op.column_dimensions["M"].width = 12  # helper column, outside the print area
LAST = 42
op.print_area = f"A1:K{LAST}"
op.page_setup.orientation = "portrait"
op.page_setup.paperSize = op.PAPERSIZE_LETTER
op.page_setup.fitToWidth = op.page_setup.fitToHeight = 1
op.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
op.page_margins.left = op.page_margins.right = op.page_margins.top = op.page_margins.bottom = 0.25
op.page_margins.header = op.page_margins.footer = 0

PKG = "Quote!$B$4"
DAYS = "Quote!$B$7"
SEC = f"Packages!$D${SEC_FIRST}:$D${SEC_LAST}"
ALSOK = f"Packages!$K${ALSO_FIRST}:$K${ALSO_LAST}"

def col(rng, letter):
    return rng.replace("$D$", f"${letter}$").replace("$K$", f"${letter}$")

# helper: number of sections in the chosen package
put(op, "M1", f"=COUNTIF(Packages!$E${SEC_FIRST}:$E${SEC_LAST},{PKG})", 8, False, FAINT)
put(op, "M2", "sections", 8, False, FAINT)

heights = {1: 6, 2: 12, 3: 10, 4: 14, 5: 46, 6: 46, 7: 28, 8: 22, 9: 8}
put(op, "B2", "368 EVENT MANAGEMENT  ×  HELLHOUND AUDIO", 7.5, True, MUTED, merge="B2:E2")
put(op, "G2", '="Prepared for "&Quote!$B$5&"  ·  Quote "&Quote!$B$8', 7.5, False, MUTED, h="right", merge="G2:J2")
put(op, "B4", '=UPPER(Quote!$B$5&"  ·  "&TEXT(Quote!$B$6,"mmmm d, yyyy"))', 8.5, True, BLUE, merge="B4:H4")
put(op, "B5", f'=IFERROR(INDEX(Packages!$B${PKG_FIRST}:$B${PKG_LAST},MATCH({PKG},Packages!$A${PKG_FIRST}:$A${PKG_LAST},0)),"")',
    36, True, wrap=True, v="center", merge="B5:H6")
put(op, "B8", "What’s in the package", 14, True, merge="B8:J8")

r = 10
price_cells = []
for i in range(1, SLOTS + 1):
    heights.update({r: 20, r + 1: 16, r + 2: 10})
    in_pkg = f"{i}<=$M$1"
    key = f'{PKG}&"|{i}"'
    sec = lambda letter: f"INDEX({col(SEC, letter)},MATCH({key},{SEC},0))"
    addon = lambda letter: f"INDEX(Quote!${letter}$18:${letter}$19,{i}-$M$1)"
    put(op, f"B{r}", f'=IFERROR(IF({in_pkg},{sec("F")},{addon("B")}&""),"")', 11, True, v="bottom", merge=f"B{r}:H{r}")
    put(op, f"B{r+1}", f'=IFERROR(IF({in_pkg},{sec("G")},{addon("C")}&""),"")', 8.5, False, MUTED, v="top",
        merge=f"B{r+1}:H{r+1}")
    put(op, f"I{r}", f'=IFERROR(IF(B{r}="","",IF({in_pkg},{sec("H")}*{DAYS},{addon("D")})),"")', 14, True,
        h="right", fmt=MONEY, merge=f"I{r}:J{r+1}")
    price_cells.append(f"I{r}")
    r += 3
assert r == 25, r

heights.update({25: 18, 26: 22, 27: 8})
put(op, "B26", f'=IF(COUNTIF({ALSOK},{PKG}&"|*")=0,"","Also included")', 14, True, merge="B26:J26")
for k in range(BULLETS):
    row, (c1, c2) = 28 + k % 3, (("B", "E"), ("G", "J"))[k // 3]
    heights[row] = 20
    put(op, f"{c1}{row}",
        f'=IFERROR("✓   "&INDEX({col(ALSOK, "M")},MATCH({PKG}&"|{k + 1}",{ALSOK},0)),"")', 9.5,
        merge=f"{c1}{row}:{c2}{row}")

heights.update({31: 30, 32: 1, 33: 1, 34: 1, 35: 1, 36: 1, 37: 1, 38: 14, 39: 42, 40: 16, 41: 26, 42: 16})
for c in range(2, 11):
    op.cell(row=38, column=c).border = Border(top=ink)
put(op, "B38", "TOTAL", 7.5, True, RED, v="bottom", merge="B38:E38")
put(op, "B39", f"=SUM({','.join(price_cells)})*(1+Quote!$B$11)", 32, True, fmt=MONEY, merge="B39:E39")
put(op, "G39", '=IF(Quote!$B$11>0,"Includes tax ("&TEXT(Quote!$B$11,"0.###%")&")","")', 8.5, False, MUTED,
    h="right", v="bottom", merge="G39:J39")
put(op, "B41", '=Quote!$B$12&"  ·  "&Quote!$B$13&"  ·  "&Quote!$B$14', 9, False, MUTED, merge="B41:J41")
put(op, "B42", '="Quote "&Quote!$B$8&"  ·  Issued "&TEXT(Quote!$B$9,"mmm d, yyyy")&"  ·  Valid until "&TEXT(Quote!$B$10,"mmm d, yyyy")',
    7.5, False, FAINT, merge="B42:J42")
for row, h in heights.items():
    op.row_dimensions[row].height = h

w = sum(op.column_dimensions[k].width * 7 + 5 for k in "ABCDEFGHIJK")
h = sum(op.row_dimensions[r].height or 15 for r in range(1, LAST + 1)) * 4 / 3
assert w <= 768 and h <= 1008, f"One-Pager is {w:.0f} x {h:.0f}px; Letter fits 768 x 1008"
print(f"One-Pager: {w:.0f} x {h:.0f}px of 768 x 1008")

wb.save(OUT)
print("wrote", OUT, [s.title for s in wb.worksheets])
