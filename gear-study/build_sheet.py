"""Build Gear_Inventory.xlsx from gear-study/data/page*.json.

Run: python3 gear-study/build_sheet.py
"""
import glob
import json
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Gear_Inventory.xlsx")
PHOTO_URL = "https://raw.githubusercontent.com/juandangooo/ok/claude/clever-lamport-wh1dmm/gear-study/photos/"

CATEGORY = {"Audio": "Audio", "Lighting": "Lighting", "Video": "Video"}
OTHER = "Staging, Rigging & Other"
TABS = ["Audio", "Lighting", "Video", OTHER]
COLORS = {"Audio": "DCEBFF", "Lighting": "FFF1C7", "Video": "E6DCFF", OTHER: "E3E3E3"}

COLUMNS = [
    ("Photo", 16), ("Name (inventory)", 34), ("Category", 13), ("Subcategory", 18),
    ("Type", 20), ("Qty owned", 9), ("Unit price (USD)", 13), ("Total value (USD)", 14),
    ("Priced as (brand / model)", 34), ("What it does", 60), ("Key specs", 36),
    ("Manual / spec sheet", 16), ("Price source", 16), ("Notes", 44), ("PDF page", 8),
]


def load():
    items = []
    for path in sorted(glob.glob(os.path.join(HERE, "data", "page*.json"))):
        page = int(os.path.basename(path)[4:6])
        for it in json.load(open(path)):
            it["page"] = page
            it["category"] = it.get("category") or CATEGORY.get(it["group"], OTHER)
            items.append(it)
    # Most expensive first; unconfirmed prices go last.
    items.sort(key=lambda i: (i["price"] is None, -(i["price"] or 0)))
    return items


def write_tab(ws, items):
    header_fill = PatternFill("solid", fgColor="1F2937")
    for c, (title, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "C2"
    wrap = Alignment(wrap_text=True, vertical="top")
    for r, it in enumerate(items, 2):
        ws.row_dimensions[r].height = 90
        fill = PatternFill("solid", fgColor=COLORS[it["category"]])
        values = [
            None, it["name"], it["category"], it["group"], it["type"], it["qty"],
            it["price"], f"=IF(ISNUMBER(G{r}),F{r}*G{r},\"\")",
            it["ref"], it["what"], it["specs"], None, None, it["notes"], it["page"],
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = wrap
        ws.cell(row=r, column=3).fill = fill
        ws.cell(row=r, column=2).font = Font(bold=True)
        for c in (7, 8):
            ws.cell(row=r, column=c).number_format = '"$"#,##0.00'
        if it["price"] is None:
            ws.cell(row=r, column=7, value="NOT CONFIRMED").font = Font(color="B91C1C", bold=True)
        for c, key, label in ((12, "manual", "Open"), (13, "price_src", "Open")):
            if it.get(key):
                cell = ws.cell(row=r, column=c, value=label)
                cell.hyperlink = it[key]
                cell.font = Font(color="1D4ED8", underline="single")
        photo = next((f"{it['id']}.{ext}" for ext in ("jpg", "png")
                      if os.path.exists(os.path.join(HERE, "photos", f"{it['id']}.{ext}"))), None)
        if photo:
            # In-cell image: stays with its row when sorting/resizing.
            ws.cell(row=r, column=1, value=f'=IMAGE("{PHOTO_URL}{photo}")')
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(items) + 1}"


def main():
    items = load()
    wb = Workbook()
    ws = wb.active
    ws.title = "All Gear"
    write_tab(ws, items)
    for tab in TABS:
        write_tab(wb.create_sheet(tab), [i for i in items if i["category"] == tab])

    info = wb.create_sheet("Read Me")
    lines = [
        "Gear Study Guide - inventory from Current RMS 'Products' export (757 products, 19 pages).",
        f"Last built: {date.today().isoformat()}",
        "",
        "Sorted by unit price, most expensive first. Items marked NOT CONFIRMED go last.",
        "Prices: new retail price of a real listing (see 'Price source'). Generic items (cables, pins) "
        "are priced using the named brand/model in 'Priced as', since the inventory doesn't list brands.",
        "Photos are in-cell IMAGE() formulas (they load from the GitHub repo and stay with their row when you sort).",
        "Category: Audio / Lighting / Video. Anything else (rigging, pipe & drape, tents, power, radios, safety) "
        f"goes under '{OTHER}'; the original inventory group is kept in 'Subcategory'.",
        "Qty 0 + 'SUB-RENTAL' = you rent it from another company when needed.",
        "A few items are moved out of the inventory's own group when it's wrong (e.g. BNC cables filed under Audio are Video); the Notes say so.",
    ]
    for r, line in enumerate(lines, 1):
        info.cell(row=r, column=1, value=line).alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 120
    wb.move_sheet(info, offset=-len(wb.sheetnames) + 1)
    wb.active = 1
    wb.save(OUT)
    priced = [i for i in items if i["price"] is not None]
    print(f"{len(items)} items, {len(priced)} priced -> {OUT}")


if __name__ == "__main__":
    main()
